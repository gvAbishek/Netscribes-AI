import logging
import asyncio
import io
from typing import List
from azure.ai.documentintelligence.aio import DocumentIntelligenceClient
from azure.ai.documentintelligence.models import AnalyzeResult, AnalyzeDocumentRequest
from azure.core.credentials import AzureKeyCredential
from azure.search.documents.aio import SearchClient
from azure.search.documents.indexes.aio import SearchIndexClient
from azure.search.documents.indexes.models import (
    SearchIndex,
    SearchField,
    SearchFieldDataType,
    SimpleField,
    SearchableField,
    VectorSearch,
    HnswAlgorithmConfiguration,
    VectorSearchProfile,
    VectorSearchAlgorithmKind,
)
from azure.search.documents.models import VectorizedQuery
from openai import AsyncAzureOpenAI
import certifi
from docx import Document as DocxDocument
from config import settings
from database import db_manager
from services.blob_service import blob_service

logger = logging.getLogger(__name__)

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file using python-docx (local, no Azure needed)."""
    doc = DocxDocument(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs.append(" | ".join(row_text))
    return "\n".join(paragraphs)


class RAGPipeline:
    def __init__(self):
        # Embedding client uses its own endpoint and key
        self.embedding_client = AsyncAzureOpenAI(
            api_key=settings.azure_openai_embedding_key,
            api_version=settings.azure_openai_api_version,
            azure_endpoint=settings.azure_openai_embedding_endpoint
        )
        self.index_name = "rag-chunks"

    async def get_embeddings(self, text: str) -> List[float]:
        try:
            response = await self.embedding_client.embeddings.create(
                input=text,
                model=settings.azure_openai_embedding_model
            )
            return response.data[0].embedding
        except Exception as e:
            logger.error(f"Embedding failed: {e}")
            raise

    def _get_di_client(self):
        endpoint = settings.azure_di_endpoint.rstrip('/')
        return DocumentIntelligenceClient(
            endpoint=endpoint,
            credential=AzureKeyCredential(settings.azure_di_key),
            connection_verify=certifi.where()
        )

    def _get_search_client(self):
        return SearchClient(
            endpoint=settings.azure_search_endpoint,
            index_name=self.index_name,
            credential=AzureKeyCredential(settings.azure_search_key),
            connection_verify=certifi.where()
        )

    def _get_index_client(self):
        return SearchIndexClient(
            endpoint=settings.azure_search_endpoint,
            credential=AzureKeyCredential(settings.azure_search_key),
            connection_verify=certifi.where()
        )

    async def search(self, query: str, user_role: str, k: int = 5):
        """Perform vector search on Azure AI Search with role context."""
        vector = await self.get_embeddings(query)
        
        async with self._get_search_client() as search_client:
            filter_str = f"allowedRoles/any(r: r eq '{user_role}')"
            
            logger.info(f"Searching index '{self.index_name}' with filter: {filter_str}")
            
            results = await search_client.search(
                search_text=query,
                vector_queries=[
                    VectorizedQuery(vector=vector, k_nearest_neighbors=k, fields="contentVector")
                ],
                filter=filter_str,
                select=["content", "title", "document_id"]
            )
            
            output = []
            async for result in results:
                output.append(result)
            return output

    async def ensure_index(self):
        """Creates the search index if it doesn't exist."""
        async with self._get_index_client() as index_client:
            try:
                await index_client.get_index(self.index_name)
                logger.info(f"Search index '{self.index_name}' already exists.")
            except Exception:
                logger.info(f"Creating search index '{self.index_name}'...")
                
                # Dimensions: 3072 for text-embedding-3-large, 1536 for small/ada-002
                dims = 3072 if "large" in settings.azure_openai_embedding_model else 1536
                
                fields = [
                    SimpleField(name="id", type=SearchFieldDataType.String, key=True),
                    SearchableField(name="content", type=SearchFieldDataType.String),
                    SimpleField(name="document_id", type=SearchFieldDataType.String, filterable=True),
                    SimpleField(name="title", type=SearchFieldDataType.String),
                    SearchField(name="allowedRoles", type=SearchFieldDataType.Collection(SearchFieldDataType.String), filterable=True),
                    SearchField(
                        name="contentVector",
                        type=SearchFieldDataType.Collection(SearchFieldDataType.Single),
                        searchable=True,
                        vector_search_dimensions=dims,
                        vector_search_profile_name="myHnswProfile"
                    )
                ]

                vector_search = VectorSearch(
                    algorithms=[
                        HnswAlgorithmConfiguration(
                            name="myHnsw",
                            kind=VectorSearchAlgorithmKind.HNSW
                        )
                    ],
                    profiles=[
                        VectorSearchProfile(
                            name="myHnswProfile",
                            algorithm_configuration_name="myHnsw"
                        )
                    ]
                )

                index = SearchIndex(
                    name=self.index_name,
                    fields=fields,
                    vector_search=vector_search
                )
                await index_client.create_index(index)
                logger.info(f"Successfully created search index '{self.index_name}'")

    async def _extract_text(self, filename: str, file_bytes: bytes) -> str:
        """Extract text: use python-docx for DOCX files, Azure DI for PDF/images."""
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        
        if ext == "docx":
            logger.info(f"Extracting text locally (python-docx): {filename}")
            return await asyncio.to_thread(extract_text_from_docx, file_bytes)
        elif ext == "txt":
            logger.info(f"Reading plain text: {filename}")
            return file_bytes.decode("utf-8", errors="replace")
        elif ext == "csv":
            logger.info(f"Reading CSV as text: {filename}")
            return file_bytes.decode("utf-8", errors="replace")
        else:
            # PDF, images, etc → use Azure Document Intelligence
            logger.info(f"Extracting text via Azure DI: {filename}")
            async with self._get_di_client() as di_client:
                poller = await di_client.begin_analyze_document(
                    "prebuilt-layout",
                    file_bytes
                )
                result = await asyncio.wait_for(poller.result(), timeout=120.0)
            return result.content

    async def process_document(self, doc_id: str):
        """
        Full pipeline: Extract text -> Chunk -> Embed -> Index
        VERSION: 7.0 (Auto Index Creation + Separate Endpoints)
        """
        logger.info(f"--- RAG PIPELINE START (V7.0): {doc_id} ---")
        
        db = db_manager.get_db()
        from bson import ObjectId
        doc = await db["rag_documents"].find_one({"_id": ObjectId(doc_id)})
        if not doc:
            logger.error(f"Document {doc_id} not found in DB")
            return

        try:
            # 0. Ensure Search Index Exists
            await self.ensure_index()

            # 1. Download blob content
            logger.info(f"Downloading blob: {doc['name']}")
            file_bytes = await asyncio.to_thread(blob_service.download_blob, doc["name"])
            logger.info(f"Downloaded {len(file_bytes)} bytes")

            # 2. Extract text (smart routing based on file type)
            content = await self._extract_text(doc["name"], file_bytes)
            logger.info(f"Text extraction complete. Content length: {len(content)}")

            if not content or not content.strip():
                raise ValueError("No text could be extracted from the document")

            # 3. Chunking
            words = content.split()
            chunks = []
            chunk_size = 500
            overlap = 100
            for i in range(0, len(words), chunk_size - overlap):
                chunk_text = " ".join(words[i : i + chunk_size])
                chunks.append(chunk_text)
                if i + chunk_size >= len(words):
                    break
            
            # 4. Embed and Index
            logger.info(f"Generating embeddings for {len(chunks)} chunks")
            search_docs = []
            for idx, chunk in enumerate(chunks):
                vector = await self.get_embeddings(chunk)
                search_docs.append({
                    "id": f"{doc_id}_{idx}",
                    "document_id": doc_id,
                    "content": chunk,
                    "title": doc["name"],
                    "allowedRoles": doc["allowedRoles"],
                    "contentVector": vector
                })
            
            async with self._get_search_client() as search_client:
                logger.info(f"Uploading {len(search_docs)} chunks to Azure Search")
                await search_client.upload_documents(documents=search_docs)
            
            # 5. Finalize
            await db["rag_documents"].update_one(
                {"_id": ObjectId(doc_id)},
                {"$set": {"status": "indexed"}}
            )
            logger.info(f"Successfully indexed: {doc['name']}")

        except Exception as e:
            import traceback
            error_msg = str(e)
            logger.error(f"Pipeline failed for {doc['name']}: {error_msg}")
            logger.error(traceback.format_exc())
            await db["rag_documents"].update_one(
                {"_id": ObjectId(doc_id)},
                {"$set": {"status": "error", "error_message": error_msg}}
            )

rag_pipeline = RAGPipeline()
