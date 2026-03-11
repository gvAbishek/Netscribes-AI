import os
import io
import pandas as pd
from docx import Document
from openpyxl import Workbook
from PyPDF2 import PdfWriter
from file_utils import extract_text_from_file

def test_extraction():
    # 1. Test TXT
    print("Testing TXT...")
    txt_content = b"Hello world, this is a text file."
    result = extract_text_from_file("test.txt", txt_content)
    print(f"Result: {result}")
    assert "Hello world" in result

    # 2. Test CSV
    print("\nTesting CSV...")
    df = pd.DataFrame({"col1": [1, 2], "col2": ["a", "b"]})
    csv_buffer = io.BytesIO()
    df.to_csv(csv_buffer, index=False)
    result = extract_text_from_file("test.csv", csv_buffer.getvalue())
    print(f"Result:\n{result}")
    assert "col1" in result and "col2" in result

    # 3. Test DOCX
    print("\nTesting DOCX...")
    doc = Document()
    doc.add_paragraph("This is a word document.")
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    result = extract_text_from_file("test.docx", doc_buffer.getvalue())
    print(f"Result: {result}")
    assert "word" in result

    # 4. Test XLSX
    print("\nTesting XLSX...")
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Excel Data"
    ws["B1"] = 123
    xls_buffer = io.BytesIO()
    wb.save(xls_buffer)
    result = extract_text_from_file("test.xlsx", xls_buffer.getvalue())
    print(f"Result: {result}")
    assert "Excel Data" in result

    # 5. Test empty file
    print("\nTesting empty file...")
    result = extract_text_from_file("empty.txt", b"")
    print(f"Result: {result}")
    assert "empty" in result

    # 6. Test unsupported extension
    print("\nTesting unsupported extension...")
    result = extract_text_from_file("test.exe", b"binary stuff")
    print(f"Result: {result}")
    assert "Unsupported" in result

    print("\nAll backend extraction tests passed!")

if __name__ == "__main__":
    test_extraction()
